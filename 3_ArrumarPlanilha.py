"""Etapa 3: valida a resposta manual do GPT e gera a planilha final do PAINEL EPT.

Execução normal:
    python 3_ArrumarPlanilha.py

Entradas principais: dados/GPT/gpt.json e dados/1_PegarDadosSiteIfc.xlsx.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import json
import os
from pathlib import Path
import re
from urllib.parse import urlsplit

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from util import agora_utc, ler_json, limpar_texto, normalizar_busca, salvar_texto_atomico


ARQUIVO_ENTRADA = Path("dados/GPT/para_gpt.json")
ARQUIVO_DADOS_OFICIAIS = Path("dados/1_PegarDadosSiteIfc.xlsx")
ARQUIVO_VOCABULARIOS = Path("dados/referencia/vocabularios.json")
ARQUIVO_SAIDA = Path("dados/dados_finais_painel_ept.xlsx")
ARQUIVO_RELATORIO = Path("dados/relatorio_validacao.txt")
ARQUIVO_RESPOSTA_GPT = Path("dados/GPT/gpt.json")
SCHEMA_VOCABULARIOS = "leme-vocabularios-v2"
SCHEMA_ENTRADA_GPT = "leme-para-classificacao-gpt-v3"
SCHEMA_RESPOSTA_GPT = "leme-classificacoes-gpt-v2"
INSTITUICAO_CORPUS = "Instituto Federal Catarinense"
CAMPUS_CORPUS = "Blumenau"
PROGRAMA_CORPUS = "ProfEPT"

LINHAS = {
    "Práticas Educativas em Educação Profissional e Tecnológica (EPT)": [
        "Práticas Educativas em EPT",
        "Práticas Educativas em Educação Profissional e Tecnológica (EPT)",
    ],
    "Organização e Memórias de Espaços Pedagógicos na Educação Profissional e Tecnológica (EPT)": [
        "Organização e Memórias de Espaços Pedagógicos na EPT",
        "Organização e Memórias de Espaços Pedagógicos na Educação Profissional e Tecnológica (EPT)",
    ],
}

MACROPROJETOS = {
    "MP1": "Propostas metodológicas e recursos didáticos em espaços formais e não formais de ensino na EPT",
    "MP2": "Inclusão e diversidade em espaços formais e não formais de ensino na EPT",
    "MP3": "Práticas Educativas no Currículo Integrado",
    "MP4": "História e memórias no contexto da EPT",
    "MP5": "Organização do currículo integrado na EPT",
    "MP6": "Organização de espaços pedagógicos na EPT",
}

CAMPOS_VOCABULARIO_ABERTO = [
    "Base Epistemológica",
    "Nível de Aplicação",
    "Tipo de Produto Educacional",
    "Área Temática",
]

CAMPOS_CLASSIFICACAO_GPT = {
    "linha_pesquisa": "Linha de Pesquisa",
    "macroprojeto": "Macroprojeto",
    "base_epistemologica": "Base Epistemológica",
    "nivel_aplicacao": "Nível de Aplicação",
    "tipo_produto": "Tipo de Produto Educacional",
    "area_tematica": "Área Temática",
    "publico_alvo": "Público-alvo",
}
CONFIANCAS_GPT = {"alta", "media", "baixa"}
COLUNAS_REVISAO_GPT = [
    "ID", "Título", "Campo", "Valor Atual", "Valor GPT", "Confiança", "Evidência", "Ação"
]

COLUNAS_FINAIS = [
    "ID",
    "Nome do Arquivo",
    "Título da Dissertação",
    "Autor",
    "Orientador Original",
    "Orientador",
    "Coorientador",
    "Ano",
    "Produto Educacional",
    "Link do PDF",
    "Link do Produto",
    "Palavras-chave",
    "Tipo de Produto Original",
    "Instituição",
    "Campus",
    "Programa",
    "Instituição/Campus",
    "Linha de Pesquisa",
    "Macroprojeto",
    "Código do Macroprojeto",
    "Classificação Encontrada Explicitamente no Texto?",
    "Trecho onde aparece Linha/Macroprojeto",
    "Confiança Linha/Macro",
    "Termos Detectados Linha/Macro",
    "Justificativa Linha/Macro",
    "Base Epistemológica Original",
    "Base Epistemológica",
    "Confiança Base Epistemológica",
    "Termos Detectados Base",
    "Nível de Aplicação Original",
    "Nível de Aplicação",
    "Tipo de Produto Educacional",
    "Área Temática Original",
    "Área Temática",
    "Público-alvo Original",
    "Público-alvo",
    "Finalidade do Produto",
    "Resumo Sintetizado",
    "Revisão Manual",
    "Observações",
    "Status do Vínculo",
    "Status da Extração",
    "SHA-256",
]

PREPOSICOES_NOMES = {"de", "da", "do", "das", "dos", "e"}
PALAVRAS_MINUSCULAS_TITULOS = {
    "a", "as", "o", "os", "um", "uma", "uns", "umas",
    "de", "da", "do", "das", "dos", "em", "no", "na", "nos", "nas",
    "ao", "aos", "à", "às", "por", "para", "com", "sem", "sob", "sobre",
    "entre", "e", "ou", "que", "como",
}
TERMOS_ESPECIAIS = {
    normalizar_busca(valor): valor
    for valor in [
        "EPT", "IFC", "IFSC", "EMI", "EJA", "EJA-EPT", "PROEJA", "ProfEPT",
        "ChatGPT", "Libras", "PPC", "PPCs", "PAE", "CISSP", "AEE", "PPI",
        "HQ", "IFS", "CEDUP", "CEDUPs", "NDB", "DUA", "TILSP", "LGBTQI",
        "NR", "PPP", "PPPs", "SC", "Moodle", "Telegram", "Kahoot", "Quizlet",
        "Wordwall", "Iramuteq", "Socrative",
    ]
}
TERMOS_ESPECIAIS.update({"ppc s": "PPCs", "chat gpt": "ChatGPT"})
CAMPOS_CATEGORIA_MULTIVALOR = {"Base Epistemológica", "Área Temática"}


def limpar_pontuacao_textual(valor: str) -> str:
    """Corrige apenas espaçamento e pontuação inequivocamente mecânicos."""
    texto = limpar_texto(valor)
    texto = re.sub(r"\s+([,.;:!?])", r"\1", texto)
    texto = re.sub(r"([,;:!?])(?=\S)", r"\1 ", texto)
    texto = re.sub(r"\s*([\u2013\u2014])\s*", r" \1 ", texto)
    texto = re.sub(r",(?:\s*,)+", ",", texto)
    texto = re.sub(r";(?:\s*;)+", ";", texto)
    texto = re.sub(r"([!?])\s*[,;]+", r"\1", texto)
    texto = re.sub(r"\(\s+", "(", texto)
    texto = re.sub(r"\s+\)", ")", texto)
    return re.sub(r"\s+", " ", texto).strip(" ,;")


def _separar_bordas_token(token: str) -> tuple[str, str, str]:
    if not re.search(r"[0-9A-Za-zÀ-ÖØ-öø-ÿ]", token):
        return "", "", token
    inicio = re.match(r"^[^0-9A-Za-zÀ-ÖØ-öø-ÿ]+", token)
    fim = re.search(r"[^0-9A-Za-zÀ-ÖØ-öø-ÿ]+$", token)
    prefixo = inicio.group(0) if inicio else ""
    sufixo = fim.group(0) if fim else ""
    limite = len(token) - len(sufixo) if sufixo else len(token)
    return prefixo, token[len(prefixo):limite], sufixo


def _capitalizar_nucleo(nucleo: str, minusculas: set[str], primeira: bool) -> str:
    especial = TERMOS_ESPECIAIS.get(normalizar_busca(nucleo))
    if especial:
        return especial
    if not nucleo:
        return nucleo
    chave = normalizar_busca(nucleo)
    if chave in minusculas and not primeira:
        return nucleo.casefold()
    partes = re.split(r"([-'’])", nucleo)
    saida = []
    indice_lexical = 0
    for parte in partes:
        if parte in {"-", "'", "’"}:
            saida.append(parte)
            continue
        letras = re.sub(r"[^A-Za-zÀ-ÖØ-öø-ÿ]", "", parte)
        if not letras:
            saida.append(parte)
            continue
        especial_parte = TERMOS_ESPECIAIS.get(normalizar_busca(parte))
        primeira_parte = primeira and indice_lexical == 0
        indice_lexical += 1
        if especial_parte:
            saida.append(especial_parte)
        elif normalizar_busca(parte) in minusculas and not primeira_parte:
            saida.append(parte.casefold())
        elif letras.isupper() or letras.islower():
            saida.append(parte[:1].upper() + parte[1:].casefold())
        else:
            saida.append(parte)
    return "".join(saida)


def padronizar_nome_pessoa(valor: str) -> str:
    original = limpar_pontuacao_textual(valor)
    if not original or original in {"---", "A revisar"}:
        return original or "---"
    tokens = original.split()
    saida = []
    for indice, token in enumerate(tokens):
        prefixo, nucleo, sufixo = _separar_bordas_token(token)
        saida.append(
            prefixo
            + _capitalizar_nucleo(nucleo, PREPOSICOES_NOMES, primeira=indice == 0)
            + sufixo
        )
    return " ".join(saida)


def padronizar_titulo(valor: str) -> str:
    texto = limpar_pontuacao_textual(valor)
    if not texto or texto in {"---", "A revisar"}:
        return texto or "---"
    texto = re.sub(r"\bchat\s+gpt\b", "ChatGPT", texto, flags=re.IGNORECASE)
    tokens = texto.split()
    saida = []
    for indice, token in enumerate(tokens):
        prefixo, nucleo, sufixo = _separar_bordas_token(token)
        primeira = indice == 0
        saida.append(
            prefixo
            + _capitalizar_nucleo(nucleo, PALAVRAS_MINUSCULAS_TITULOS, primeira)
            + sufixo
        )
    return " ".join(saida)


def partes_multivalor(valor: str) -> list[str]:
    """Separa listas explícitas; barras internas sem espaços permanecem intactas."""
    return [
        limpar_texto(parte)
        for parte in re.split(r"\s+/\s+|\s*;\s*", limpar_texto(valor))
        if limpar_texto(parte)
    ]


def argumentos() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--classificacoes-gpt",
        type=Path,
        default=ARQUIVO_RESPOSTA_GPT,
        help=f"resposta única do GPT (padrão: {ARQUIVO_RESPOSTA_GPT})",
    )
    return parser.parse_args()




def indice_alias(mapeamento: dict[str, list[str]]) -> dict[str, str]:
    return {
        normalizar_busca(alias): canonico
        for canonico, aliases in mapeamento.items()
        for alias in set(aliases + [canonico])
    }


def indice_canonicos(canonicos: dict) -> dict[str, tuple[str, dict]]:
    """Indexa aliases declarados no vocabulário, incluindo o valor canônico."""
    indice: dict[str, tuple[str, dict]] = {}
    for canonico, configuracao in canonicos.items():
        if isinstance(configuracao, list):
            configuracao = {"aliases": configuracao, "confianca": "alta", "regra": "alias controlado"}
        for alias in set(configuracao.get("aliases", []) + [canonico]):
            indice[normalizar_busca(alias)] = (canonico, configuracao)
    return indice


def indice_atomico_categoria(
    campo: str, vocabulario: dict
) -> tuple[dict[str, str], dict[str, int]]:
    """Indexa cada componente canônico de categorias explicitamente multivalor."""
    canonicos = vocabulario.get("categorias", {}).get(campo, {}).get("canonicos", {})
    indice: dict[str, str] = {}
    ordem: dict[str, int] = {}
    for canonico, configuracao in canonicos.items():
        partes_canonicas = partes_multivalor(canonico)
        for parte in partes_canonicas:
            chave = normalizar_busca(parte)
            if chave not in indice:
                indice[chave] = parte
                ordem[chave] = len(ordem)
        aliases = configuracao if isinstance(configuracao, list) else configuracao.get("aliases", [])
        for alias in aliases:
            partes_alias = partes_multivalor(alias)
            if len(partes_alias) != len(partes_canonicas):
                continue
            for parte_alias, parte_canonica in zip(partes_alias, partes_canonicas):
                indice.setdefault(
                    normalizar_busca(parte_alias),
                    indice[normalizar_busca(parte_canonica)],
                )
    return indice, ordem


def normalizar_categoria_multivalor(
    campo: str, valor: str, vocabulario: dict
) -> tuple[str, str | None]:
    texto = limpar_texto(valor)
    if not texto or normalizar_busca(texto) == "a revisar":
        return "A revisar", None
    indice, ordem = indice_atomico_categoria(campo, vocabulario)
    partes = []
    desconhecidas = []
    vistos = set()
    for parte_original in partes_multivalor(texto):
        chave_original = normalizar_busca(parte_original)
        canonico = indice.get(chave_original, parte_original)
        chave_canonica = normalizar_busca(canonico)
        if chave_original not in indice:
            desconhecidas.append(parte_original)
        if chave_canonica not in vistos:
            vistos.add(chave_canonica)
            partes.append(canonico)
    partes.sort(
        key=lambda item: (
            ordem.get(normalizar_busca(item), len(ordem)),
            normalizar_busca(item),
        )
    )
    erro = None
    if desconhecidas:
        erro = f"{campo} possui categoria sem alias confirmado: {', '.join(desconhecidas)}"
    return " / ".join(partes) or "A revisar", erro


def chave_pessoa(valor: str) -> str:
    """Compara nomes sem título acadêmico, caixa ou pontuação."""
    texto = limpar_texto(valor)
    texto = re.sub(
        r"^\s*(?:(?:prof(?:essor|essora)?|prof[ªa]?|dr[ªa]?|dra)\.?\s*[:.-]?\s*)+",
        "",
        texto,
        flags=re.IGNORECASE,
    )
    texto = texto.strip(" .:;,-")
    return normalizar_busca(texto)


def normalizar_orientador(valor: str, vocabulario: dict) -> tuple[str, str | None, str]:
    original = limpar_texto(valor)
    if not original:
        return "---", "orientador vazio", "sem valor"
    indice = {}
    for canonico, configuracao in vocabulario.get("orientadores", {}).get("canonicos", {}).items():
        for alias in set(configuracao.get("aliases", []) + [canonico]):
            indice[chave_pessoa(alias)] = (canonico, configuracao)
    encontrado = indice.get(chave_pessoa(original))
    if not encontrado:
        nome_sem_titulo = re.sub(
            r"^\s*(?:(?:prof(?:essor|essora)?|prof[ªa]?|dr[ªa]?|dra)\.?\s*[:.-]?\s*)+",
            "",
            original,
            flags=re.IGNORECASE,
        ).strip(" .:;,-")
        nome_formatado = padronizar_nome_pessoa(nome_sem_titulo)
        return nome_formatado or "---", f"orientador sem alias confirmado: {original}", "remoção de título não confirmado"
    canonico, configuracao = encontrado
    return padronizar_nome_pessoa(canonico), None, configuracao.get("regra", "alias de pessoa confirmado")


def normalizar_linha(valor: str, vocabulario: dict) -> tuple[str, str | None]:
    texto = limpar_texto(valor)
    if not texto or texto == "A revisar":
        return "A revisar", None
    canonico = indice_alias(vocabulario["linhas_de_pesquisa"]).get(normalizar_busca(texto))
    if not canonico:
        return "A revisar", f"valor não reconhecido preservado para revisão: {texto}"
    return canonico, None


def normalizar_macro(valor: str, codigo: str, vocabulario: dict) -> tuple[str, str, str | None]:
    texto = limpar_texto(valor)
    codigo_texto = limpar_texto(codigo).upper()
    macros = vocabulario["macroprojetos"]
    codigo_no_texto = re.match(r"^(MP[1-6])\s*[–—:-]?\s*(.*)$", texto, flags=re.IGNORECASE)
    if codigo_no_texto:
        codigo_detectado = codigo_no_texto.group(1).upper()
        descricao = limpar_texto(codigo_no_texto.group(2))
    else:
        codigo_detectado = ""
        descricao = texto
    por_descricao = {normalizar_busca(descricao_oficial): chave for chave, descricao_oficial in macros.items()}
    codigo_descricao = por_descricao.get(normalizar_busca(descricao)) if descricao else ""
    candidatos = {item for item in [codigo_texto, codigo_detectado, codigo_descricao] if item}
    if not texto or texto == "A revisar":
        return "A revisar", "A revisar", None
    if len(candidatos) != 1 or next(iter(candidatos), "") not in macros:
        detalhe = f"macroprojeto/código inconsistente: macro={texto!r}, código={codigo!r}"
        return "A revisar", "A revisar", detalhe
    codigo_final = next(iter(candidatos))
    return f"{codigo_final} – {macros[codigo_final]}", codigo_final, None


def normalizar_valor_aberto(campo: str, valor: str, vocabulario: dict) -> tuple[str, str | None]:
    texto = limpar_texto(valor)
    if not texto or texto == "A revisar":
        return "A revisar", None
    if campo in CAMPOS_CATEGORIA_MULTIVALOR:
        return normalizar_categoria_multivalor(campo, texto, vocabulario)
    categoria = vocabulario.get("categorias", {}).get(campo, {})
    encontrado = indice_canonicos(categoria.get("canonicos", {})).get(normalizar_busca(texto))
    if encontrado:
        return encontrado[0], None

    partes = partes_multivalor(texto)
    if len(partes) > 1:
        indice = indice_canonicos(categoria.get("canonicos", {}))
        canonicos = []
        desconhecidos = []
        vistos = set()
        for parte in partes:
            encontrado_parte = indice.get(normalizar_busca(parte))
            if not encontrado_parte:
                desconhecidos.append(parte)
                continue
            canonico = encontrado_parte[0]
            chave = normalizar_busca(canonico)
            if chave not in vistos:
                vistos.add(chave)
                canonicos.append(canonico)
        if not desconhecidos and canonicos:
            ordem = {normalizar_busca(item): pos for pos, item in enumerate(categoria.get("canonicos", {}))}
            canonicos.sort(key=lambda item: (ordem.get(normalizar_busca(item), len(ordem)), normalizar_busca(item)))
            return " / ".join(canonicos), None
    # Compatibilidade com vocabulários antigos usados em testes unitários.
    aceitos = vocabulario.get("valores_historicos_aceitos", {}).get(campo, [])
    indice_antigo = {normalizar_busca(item): item for item in aceitos}
    if normalizar_busca(texto) in indice_antigo:
        return indice_antigo[normalizar_busca(texto)], None
    if categoria.get("preservar_nao_mapeados", False):
        pendentes = {
            normalizar_busca(item.get("valor") if isinstance(item, dict) else item)
            for item in categoria.get("pendentes_revisao", [])
        }
        erro = f"{campo} preservado para revisão de equivalência: {texto}" if normalizar_busca(texto) in pendentes else None
        return texto, erro
    return "A revisar", f"{campo} fora do vocabulário: {texto}"


def normalizar_publico_alvo(valor: str, vocabulario: dict) -> tuple[str, str | None, str]:
    original = limpar_texto(valor)
    if not original or original in {"---", "A revisar"}:
        return "A revisar", "público-alvo ausente", "campo ausente"
    regras = vocabulario.get("publico_alvo", {}).get("regras_palavras_chave", [])
    chave = normalizar_busca(original)
    categorias = []
    regras_usadas = []
    for regra in regras:
        if any(normalizar_busca(termo) in chave for termo in regra.get("termos", [])):
            categoria = regra["canonico"]
            if categoria not in categorias:
                categorias.append(categoria)
                regras_usadas.append(regra.get("regra", "palavra-chave controlada"))
    if not categorias:
        return "A revisar", f"público-alvo sem regra segura: {original}", "sem correspondência"
    ordem = vocabulario.get("publico_alvo", {}).get("ordem", [])
    categorias.sort(key=lambda item: ordem.index(item) if item in ordem else len(ordem))
    return "; ".join(categorias), None, "; ".join(dict.fromkeys(regras_usadas))


def url_valida(valor: str) -> bool:
    if not valor:
        return False
    partes = urlsplit(valor)
    return partes.scheme in {"http", "https"} and bool(partes.netloc)


def preferir(*valores) -> str:
    for valor in valores:
        texto = limpar_texto(valor)
        if texto:
            return texto
    return ""


def construir_lexico_textual(registros: list[dict]) -> set[str]:
    """Reúne palavras observadas separadamente para detectar possíveis aglutinações."""
    lexico = set()
    for registro in registros:
        for campo in ["titulo", "produto_educacional"]:
            for palavra in re.findall(
                r"[A-Za-zÀ-ÖØ-öø-ÿ]{4,}",
                limpar_texto(registro.get(campo)),
            ):
                lexico.add(normalizar_busca(palavra))
    return lexico


def palavras_possivelmente_coladas(valor: str, lexico: set[str] | None) -> list[str]:
    if not lexico:
        return []
    suspeitas = []
    for token in re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ]{12,}", limpar_texto(valor)):
        chave = normalizar_busca(token)
        if chave in TERMOS_ESPECIAIS:
            continue
        for posicao in range(5, len(chave) - 4):
            esquerda, direita = chave[:posicao], chave[posicao:]
            if esquerda in {
                "trans", "inter", "intra", "extra", "auto", "multi", "micro",
                "macro", "sobre", "contra", "super", "hiper",
            }:
                continue
            if esquerda in lexico and direita in lexico:
                suspeitas.append(f"{token} (possível separação: {esquerda} + {direita})")
                break
    return suspeitas


def detectar_problemas_textuais(
    registro: dict,
    autor: str,
    orientador: str,
    coorientador: str,
    titulo: str,
    produto: str,
    lexico: set[str] | None = None,
) -> list[str]:
    problemas = []
    pessoas = {
        "autor": autor,
        "orientador": orientador,
        "coorientador": coorientador,
    }
    titulo_chave = normalizar_busca(titulo)
    produto_chave = normalizar_busca(produto)
    for rotulo, pessoa in pessoas.items():
        chave = normalizar_busca(pessoa)
        if not chave or pessoa == "---":
            continue
        if titulo_chave == chave:
            problemas.append(f"campo possivelmente deslocado: título igual ao {rotulo}")
        if produto_chave == chave:
            problemas.append(f"campo possivelmente deslocado: produto igual ao {rotulo}")

    autor_original = limpar_texto(registro.get("autor"))
    if re.search(
        r"(?:^|\s)[B-DF-HJ-NP-TV-ZÇ]\s+[A-ZÀ-ÖØ-Þ]{3,}(?=\s|$)",
        autor_original,
    ):
        problemas.append(f"espaço suspeito no nome do autor: {autor_original}")
    if any(sinal in autor_original for sinal in [":", ";", "?"]):
        problemas.append(f"pontuação atípica no nome do autor: {autor_original}")

    for campo, valor in [
        ("título", registro.get("titulo")),
        ("produto", registro.get("produto_educacional")),
    ]:
        for suspeita in palavras_possivelmente_coladas(limpar_texto(valor), lexico):
            problemas.append(f"palavra possivelmente colada no {campo}: {suspeita}")
    return list(dict.fromkeys(problemas))


def construir_linha(
    registro: dict,
    vocabulario: dict,
    lexico: set[str] | None = None,
) -> tuple[dict, list[dict], list[str]]:
    identificador = limpar_texto(registro.get("id"))
    historico = registro.get("classificacoes", {}).get("valores", {})
    originais_historicos = registro.get("classificacoes", {}).get("valores_originais", {})
    conteudo = registro.get("conteudo_pdf", {})
    alteracoes: list[dict] = []
    pendencias: list[str] = []

    autor_original = preferir(registro.get("autor"), "---")
    autor = padronizar_nome_pessoa(autor_original)
    coorientador_original = preferir(registro.get("coorientador"), "---")
    coorientador = padronizar_nome_pessoa(coorientador_original)
    titulo_original = preferir(registro.get("titulo"), "---")
    titulo = padronizar_titulo(titulo_original)
    produto_original = preferir(registro.get("produto_educacional"), "---")
    produto = padronizar_titulo(produto_original)
    for campo, original, padronizado, regra in [
        ("Autor", autor_original, autor, "capitalização nominal conservadora"),
        ("Coorientador", coorientador_original, coorientador, "capitalização nominal conservadora"),
        ("Título da Dissertação", titulo_original, titulo, "capitalização e pontuação conservadoras"),
        ("Produto Educacional", produto_original, produto, "capitalização e pontuação conservadoras"),
    ]:
        if padronizado != original:
            alteracoes.append(
                {
                    "ID": identificador,
                    "Campo": campo,
                    "Original": original,
                    "Padronizado": padronizado,
                    "Regra": regra,
                }
            )

    orientador_original = preferir(registro.get("orientador"), "---")
    orientador, erro, regra_orientador = normalizar_orientador(orientador_original, vocabulario)
    if erro:
        pendencias.append(erro)
    if orientador != orientador_original:
        alteracoes.append(
            {
                "ID": identificador,
                "Campo": "Orientador",
                "Original": orientador_original,
                "Padronizado": orientador,
                "Regra": regra_orientador,
            }
        )

    linha_original = preferir(historico.get("Linha de Pesquisa"), "A revisar")
    linha_original_auditoria = preferir(
        originais_historicos.get("Linha de Pesquisa"), linha_original
    )
    linha, erro = normalizar_linha(linha_original, vocabulario)
    if erro:
        pendencias.append(erro)
    if linha != linha_original_auditoria:
        alteracoes.append({"ID": identificador, "Campo": "Linha de Pesquisa", "Original": linha_original_auditoria, "Padronizado": linha, "Regra": "alias controlado"})

    macro_original = preferir(historico.get("Macroprojeto"), "A revisar")
    macro_original_auditoria = preferir(
        originais_historicos.get("Macroprojeto"), macro_original
    )
    codigo_original = preferir(historico.get("Código do Macroprojeto"), "A revisar")
    macro, codigo, erro = normalizar_macro(macro_original, codigo_original, vocabulario)
    if erro:
        pendencias.append(erro)
    if macro != macro_original_auditoria:
        alteracoes.append({"ID": identificador, "Campo": "Macroprojeto", "Original": macro_original_auditoria, "Padronizado": macro, "Regra": "código e descrição controlados"})
    if codigo != codigo_original:
        alteracoes.append({"ID": identificador, "Campo": "Código do Macroprojeto", "Original": codigo_original, "Padronizado": codigo, "Regra": "código derivado do vocabulário"})

    abertos = {}
    abertos_originais = {}
    for campo in CAMPOS_VOCABULARIO_ABERTO:
        original = preferir(historico.get(campo), "A revisar")
        abertos_originais[campo] = preferir(originais_historicos.get(campo), original)
        abertos[campo], erro = normalizar_valor_aberto(campo, original, vocabulario)
        if erro:
            pendencias.append(erro)
        if abertos[campo] != abertos_originais[campo]:
            alteracoes.append({"ID": identificador, "Campo": campo, "Original": abertos_originais[campo], "Padronizado": abertos[campo], "Regra": "vocabulário controlado"})

    publico_fonte = preferir(historico.get("Público-alvo"), conteudo.get("publico_alvo_extraido"), "A revisar")
    publico_original = preferir(originais_historicos.get("Público-alvo"), publico_fonte)
    publico, erro, regra_publico = normalizar_publico_alvo(publico_fonte, vocabulario)
    if erro:
        pendencias.append(erro)
    if publico != publico_original:
        alteracoes.append(
            {
                "ID": identificador,
                "Campo": "Público-alvo",
                "Original": publico_original,
                "Padronizado": publico,
                "Regra": regra_publico,
            }
        )

    pendencias.extend(
        detectar_problemas_textuais(
            registro,
            autor,
            orientador,
            coorientador,
            titulo,
            produto,
            lexico,
        )
    )

    conflitos = registro.get("validacao", {}).get("conflitos", [])
    vinculo = bool(registro.get("validacao", {}).get("vinculo_confirmado"))
    status_extracao = limpar_texto(registro.get("extracao_pdf", {}).get("status"))
    if not vinculo:
        pendencias.append("vínculo entre trabalho e PDF não confirmado")
    if status_extracao != "sucesso":
        pendencias.append(f"extração do PDF: {status_extracao or 'sem status'}")
    if conflitos:
        pendencias.extend(f"conflito: {item.get('campo', 'não especificado')}" for item in conflitos)
    if any(valor == "A revisar" for valor in [linha, macro, codigo, *abertos.values()]):
        pendencias.append("uma ou mais classificações aguardam revisão")

    observacoes = preferir(historico.get("Observações"))
    if pendencias:
        complemento = "Pendências automáticas: " + "; ".join(dict.fromkeys(pendencias))
        observacoes = f"{observacoes} | {complemento}" if observacoes else complemento

    instituicao = preferir(registro.get("instituicao"), INSTITUICAO_CORPUS)
    campus = preferir(registro.get("campus"), CAMPUS_CORPUS)
    programa = preferir(registro.get("programa"), PROGRAMA_CORPUS)
    instituicao_campus = preferir(registro.get("instituicao_campus"), f"IFC {campus}")

    linha_final = {
        "ID": identificador,
        "Nome do Arquivo": Path(registro.get("arquivo_pdf", {}).get("caminho") or "").name,
        "Título da Dissertação": titulo,
        "Autor": autor,
        "Orientador Original": orientador_original,
        "Orientador": orientador,
        "Coorientador": coorientador,
        "Ano": preferir(registro.get("ano"), "---"),
        "Produto Educacional": produto,
        "Link do PDF": preferir(registro.get("url_pdf")),
        "Link do Produto": preferir(registro.get("url_produto")),
        "Palavras-chave": "; ".join(
            texto for valor in (conteudo.get("palavras_chave") or []) if (texto := limpar_texto(valor))
        ),
        "Tipo de Produto Original": preferir(
            originais_historicos.get("Tipo de Produto Educacional"),
            historico.get("Tipo de Produto Educacional"),
            "A revisar",
        ),
        "Instituição": instituicao,
        "Campus": campus,
        "Programa": programa,
        "Instituição/Campus": instituicao_campus,
        "Linha de Pesquisa": linha,
        "Macroprojeto": macro,
        "Código do Macroprojeto": codigo,
        "Classificação Encontrada Explicitamente no Texto?": preferir(historico.get("Classificação Encontrada Explicitamente no Texto?"), "Não"),
        "Trecho onde aparece Linha/Macroprojeto": preferir(historico.get("Trecho onde aparece Linha/Macroprojeto")),
        "Confiança Linha/Macro": preferir(historico.get("Confiança Linha/Macro"), "A revisar"),
        "Termos Detectados Linha/Macro": preferir(historico.get("Termos Detectados Linha/Macro")),
        "Justificativa Linha/Macro": preferir(historico.get("Justificativa Linha/Macro")),
        "Base Epistemológica Original": abertos_originais["Base Epistemológica"],
        **abertos,
        "Confiança Base Epistemológica": preferir(historico.get("Confiança Base Epistemológica"), "A revisar"),
        "Termos Detectados Base": preferir(historico.get("Termos Detectados Base")),
        "Nível de Aplicação Original": abertos_originais["Nível de Aplicação"],
        "Área Temática Original": abertos_originais["Área Temática"],
        "Público-alvo Original": publico_original,
        "Público-alvo": publico,
        "Finalidade do Produto": preferir(historico.get("Finalidade do Produto"), conteudo.get("produto_educacional_extraido"), "---"),
        "Resumo Sintetizado": preferir(historico.get("Resumo Sintetizado"), conteudo.get("resumo"), "---"),
        "Revisão Manual": "Sim" if pendencias else preferir(historico.get("Revisão Manual"), "Não"),
        "Observações": observacoes,
        "Status do Vínculo": "Confirmado" if vinculo else "A revisar",
        "Status da Extração": status_extracao or "A revisar",
        "SHA-256": preferir(registro.get("arquivo_pdf", {}).get("sha256"), registro.get("extracao_pdf", {}).get("sha256_validado")),
    }
    return {coluna: linha_final.get(coluna, "") for coluna in COLUNAS_FINAIS}, alteracoes, pendencias


def validar_conjunto(linhas: list[dict]) -> tuple[list[dict], list[dict], dict]:
    validacoes: list[dict] = []
    revisoes: list[dict] = []
    ids = [linha["ID"] for linha in linhas]
    contagem_ids = Counter(ids)
    urls_pdf: defaultdict[str, list[str]] = defaultdict(list)
    autores_titulos: defaultdict[tuple[str, str], list[str]] = defaultdict(list)
    for linha in linhas:
        identificador = linha["ID"]
        problemas = []
        if not re.fullmatch(r"EPT-\d{4,}", identificador):
            problemas.append("ID inválido")
        if contagem_ids[identificador] > 1:
            problemas.append("ID duplicado")
        for campo in ["Ano", "Autor", "Orientador", "Título da Dissertação", "Produto Educacional"]:
            if linha[campo] in {"", "---"}:
                problemas.append(f"campo essencial vazio: {campo}")
        if linha["Instituição"] != INSTITUICAO_CORPUS:
            problemas.append("instituição divergente do corpus ProfEPT/IFC")
        if linha["Campus"] in {"", "---", "A revisar"}:
            problemas.append("campus ausente ou não confirmado")
        if linha["Programa"] != PROGRAMA_CORPUS:
            problemas.append("programa divergente do corpus ProfEPT/IFC")
        if "profept" in normalizar_busca(linha["Instituição"]):
            problemas.append("ProfEPT usado incorretamente como instituição")
        for campo in ["Link do PDF", "Link do Produto"]:
            if linha[campo] and not url_valida(linha[campo]):
                problemas.append(f"URL inválida: {campo}")
        if linha["Link do PDF"]:
            urls_pdf[linha["Link do PDF"]].append(identificador)
        autores_titulos[(normalizar_busca(linha["Autor"]), normalizar_busca(linha["Título da Dissertação"]))].append(identificador)
        if linha["Revisão Manual"] == "Sim":
            revisoes.append({"ID": identificador, "Motivos": linha["Observações"]})
        validacoes.append({"ID": identificador, "Status": "ERRO" if problemas else "OK", "Problemas": "; ".join(problemas)})

    duplicidades = []
    for url, registros in urls_pdf.items():
        if len(registros) > 1:
            duplicidades.append(f"URL de PDF repetida em {', '.join(registros)}: {url}")
    for chave, registros in autores_titulos.items():
        if chave != ("", "") and len(registros) > 1:
            duplicidades.append(f"autor/título repetidos em {', '.join(registros)}")
    resumo = {
        "quantidade": len(linhas),
        "ids_unicos": len(set(ids)),
        "erros_estruturais": sum(item["Status"] == "ERRO" for item in validacoes),
        "revisoes_manuais": len(revisoes),
        "duplicidades": duplicidades,
        "vinculos_confirmados": sum(linha["Status do Vínculo"] == "Confirmado" for linha in linhas),
        "extracoes_sucesso": sum(linha["Status da Extração"] == "sucesso" for linha in linhas),
    }
    return validacoes, revisoes, resumo


def formatar_aba_planilha(planilha) -> None:
    """Aplica apresentação consistente sem alterar valores ou estrutura de dados."""
    preenchimento = PatternFill("solid", fgColor="4D5AFF")
    fonte = Font(color="FFFFFF", bold=True)
    alinhamento_cabecalho = Alignment(horizontal="center", vertical="center", wrap_text=True)
    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = planilha.dimensions
    planilha.row_dimensions[1].height = 34

    cabecalhos = {celula.column: limpar_texto(celula.value) for celula in planilha[1]}
    longos = {
        "Título da Dissertação", "Produto Educacional", "Resumo Sintetizado",
        "Finalidade do Produto", "Observações", "Problemas", "Motivos", "Evidência",
        "Valor Atual", "Valor GPT", "Original", "Padronizado", "Regra",
        "Trecho onde aparece Linha/Macroprojeto", "Justificativa Linha/Macro",
    }
    largos = {
        "Autor", "Orientador Original", "Orientador", "Coorientador",
        "Base Epistemológica", "Base Epistemológica Original", "Área Temática",
        "Área Temática Original", "Linha de Pesquisa", "Macroprojeto",
        "Público-alvo", "Público-alvo Original", "Nome canônico", "Alias encontrado",
    }
    estreitos = {"ID", "Ano", "Status", "Ação", "Confiança", "Ocorrências"}
    for celula in planilha[1]:
        celula.fill = preenchimento
        celula.font = fonte
        celula.alignment = alinhamento_cabecalho

    for indice_coluna, cabecalho in cabecalhos.items():
        valores = [
            limpar_texto(planilha.cell(linha, indice_coluna).value)
            for linha in range(1, min(planilha.max_row, 80) + 1)
        ]
        largura_conteudo = max((len(valor) for valor in valores), default=10) + 2
        if cabecalho in longos:
            largura = min(max(largura_conteudo, 34), 64)
        elif cabecalho in largos:
            largura = min(max(largura_conteudo, 24), 42)
        elif cabecalho in estreitos:
            largura = min(max(largura_conteudo, 10), 18)
        elif "Link" in cabecalho or "URL" in cabecalho:
            largura = 38
        elif cabecalho == "SHA-256":
            largura = 22
        else:
            largura = min(max(largura_conteudo, 12), 28)
        planilha.column_dimensions[get_column_letter(indice_coluna)].width = largura
        quebrar = cabecalho in longos or cabecalho in largos
        for linha in range(2, planilha.max_row + 1):
            planilha.cell(linha, indice_coluna).alignment = Alignment(
                vertical="top", wrap_text=quebrar
            )


def salvar_planilha(
    linhas: list[dict],
    alteracoes: list[dict],
    validacoes: list[dict],
    revisoes: list[dict],
    aliases_orientador: list[dict] | None = None,
    revisao_gpt: list[dict] | None = None,
) -> None:
    ARQUIVO_SAIDA.parent.mkdir(parents=True, exist_ok=True)
    temporario = ARQUIVO_SAIDA.with_name(ARQUIVO_SAIDA.stem + ".tmp.xlsx")
    try:
        with pd.ExcelWriter(temporario, engine="openpyxl") as escritor:
            pd.DataFrame(linhas, columns=COLUNAS_FINAIS).to_excel(escritor, sheet_name="Planilha1", index=False)
            pd.DataFrame(alteracoes, columns=["ID", "Campo", "Original", "Padronizado", "Regra"]).to_excel(escritor, sheet_name="Alterações", index=False)
            pd.DataFrame(validacoes, columns=["ID", "Status", "Problemas"]).to_excel(escritor, sheet_name="Validação", index=False)
            pd.DataFrame(revisoes, columns=["ID", "Motivos"]).to_excel(escritor, sheet_name="Revisão Manual", index=False)
            pd.DataFrame(
                aliases_orientador or [],
                columns=["Nome canônico", "Alias encontrado", "Ocorrências", "Confiança", "Regra", "Status"],
            ).to_excel(escritor, sheet_name="Aliases Orientador", index=False)
            pd.DataFrame(
                revisao_gpt or [], columns=COLUNAS_REVISAO_GPT
            ).to_excel(escritor, sheet_name="Revisão GPT", index=False)
            for planilha in escritor.book.worksheets:
                formatar_aba_planilha(planilha)
        os.replace(temporario, ARQUIVO_SAIDA)
    finally:
        temporario.unlink(missing_ok=True)


def inventario_antes(registros: list[dict]) -> dict[str, Counter]:
    campos = ["Linha de Pesquisa", "Macroprojeto", *CAMPOS_VOCABULARIO_ABERTO, "Público-alvo"]
    inventario = {campo: Counter() for campo in campos}
    inventario["Orientador"] = Counter()
    for registro in registros:
        historico = registro.get("classificacoes", {}).get("valores", {})
        originais = registro.get("classificacoes", {}).get("valores_originais", {})
        conteudo = registro.get("conteudo_pdf", {})
        inventario["Orientador"][preferir(registro.get("orientador"), "A revisar")] += 1
        for campo in campos:
            if campo == "Público-alvo":
                valor = preferir(
                    originais.get(campo), historico.get(campo),
                    conteudo.get("publico_alvo_extraido"), "A revisar"
                )
            else:
                valor = preferir(originais.get(campo), historico.get(campo), "A revisar")
            inventario[campo][valor] += 1
    return inventario


def inventario_depois(linhas: list[dict]) -> dict[str, Counter]:
    campos = ["Orientador", "Linha de Pesquisa", "Macroprojeto", *CAMPOS_VOCABULARIO_ABERTO, "Público-alvo"]
    return {campo: Counter(linha[campo] for linha in linhas) for campo in campos}


def tabela_aliases_orientador(registros: list[dict], vocabulario: dict) -> list[dict]:
    contagem = Counter(preferir(item.get("orientador"), "A revisar") for item in registros)
    saida = []
    canonicos = vocabulario.get("orientadores", {}).get("canonicos", {})
    indice_config = {
        chave_pessoa(alias): (canonico, config)
        for canonico, config in canonicos.items()
        for alias in set(config.get("aliases", []) + [canonico])
    }
    for alias, quantidade in sorted(contagem.items(), key=lambda item: (-item[1], normalizar_busca(item[0]))):
        canonico, erro, regra = normalizar_orientador(alias, vocabulario)
        config = indice_config.get(chave_pessoa(alias), ({}, {}))[1]
        saida.append(
            {
                "Nome canônico": canonico,
                "Alias encontrado": alias,
                "Ocorrências": quantidade,
                "Confiança": config.get("confianca", "baixa" if erro else "alta"),
                "Regra": regra,
                "Status": "Revisar" if erro else "Confirmado",
            }
        )
    for pendente in vocabulario.get("orientadores", {}).get("possiveis_equivalencias_nao_aplicadas", []):
        saida.append(
            {
                "Nome canônico": pendente.get("nome_a", ""),
                "Alias encontrado": pendente.get("nome_b", ""),
                "Ocorrências": "",
                "Confiança": "baixa",
                "Regra": pendente.get("motivo", "sem evidência suficiente"),
                "Status": "Não agrupado; revisar",
            }
        )
    return saida


def limitar_contexto_gpt(valor, limite: int) -> str | None:
    texto = limpar_texto(valor)
    if not texto:
        return None
    if len(texto) <= limite:
        return texto
    return texto[:limite].rsplit(" ", 1)[0] + " […]"






def valores_permitidos_gpt(linhas: list[dict], vocabulario: dict) -> dict:
    permitidos = {}
    for chave, coluna in CAMPOS_CLASSIFICACAO_GPT.items():
        if chave == "linha_pesquisa":
            valores = set(vocabulario["linhas_de_pesquisa"])
        elif chave == "macroprojeto":
            valores = {
                f"{codigo} – {descricao}"
                for codigo, descricao in vocabulario["macroprojetos"].items()
            }
        elif chave == "publico_alvo":
            valores = set(vocabulario.get("publico_alvo", {}).get("ordem", []))
        else:
            valores = {linha[coluna] for linha in linhas if linha[coluna] != "A revisar"}
            valores.update(
                vocabulario.get("categorias", {}).get(coluna, {}).get("canonicos", {})
            )
        valores.discard("A revisar")
        permitidos[chave] = sorted(valores, key=normalizar_busca) + ["A revisar"]
    return permitidos












def caminhos_resposta_gpt(caminho: Path) -> list[Path]:
    if caminho.is_file():
        return [caminho]
    if caminho.is_dir():
        raise ValueError(
            f"A resposta deve ser um único arquivo JSON, não uma pasta: {caminho}"
        )
    raise FileNotFoundError(f"Resposta GPT não encontrada: {caminho}")


def canonizar_valor_gpt(
    chave: str, valor, vocabulario: dict, permitidos: dict[str, list[str]]
) -> str:
    texto = limpar_texto(valor)
    if not texto:
        raise ValueError(f"{chave}: valor vazio")
    if normalizar_busca(texto) == "a revisar":
        return "A revisar"
    if chave == "linha_pesquisa":
        canonico, erro = normalizar_linha(texto, vocabulario)
        if erro:
            raise ValueError(f"linha_pesquisa inválida: {texto}")
        return canonico
    if chave == "macroprojeto":
        canonico, _, erro = normalizar_macro(texto, "", vocabulario)
        if erro:
            raise ValueError(f"macroprojeto inválido: {texto}")
        return canonico
    if chave == "publico_alvo":
        indice = {
            normalizar_busca(item): item
            for item in vocabulario.get("publico_alvo", {}).get("ordem", [])
        }
        partes = [limpar_texto(item) for item in texto.split(";") if limpar_texto(item)]
        if not partes:
            raise ValueError("publico_alvo vazio")
        canonicos = []
        for parte in partes:
            canonico = indice.get(normalizar_busca(parte))
            if not canonico:
                raise ValueError(f"publico_alvo inválido: {parte}")
            if canonico not in canonicos:
                canonicos.append(canonico)
        ordem = vocabulario.get("publico_alvo", {}).get("ordem", [])
        canonicos.sort(key=lambda item: ordem.index(item))
        return "; ".join(canonicos)
    indice = {
        normalizar_busca(item): item
        for item in permitidos.get(chave, [])
        if item != "A revisar"
    }
    coluna = CAMPOS_CLASSIFICACAO_GPT[chave]
    normalizado, _ = normalizar_valor_aberto(coluna, texto, vocabulario)
    canonico = indice.get(normalizar_busca(normalizado))
    if not canonico:
        raise ValueError(f"{chave} inválido ou ainda não aprovado no vocabulário: {texto}")
    return canonico


def validar_respostas_gpt(
    caminho: Path,
    registros: list[dict],
    linhas: list[dict],
    vocabulario: dict,
) -> dict:
    esperados = {registro["id"] for registro in registros}
    linhas_por_id = {linha["ID"]: linha for linha in linhas}
    permitidos = valores_permitidos_gpt(linhas, vocabulario)
    sugestoes = {}
    avisos = []
    vocabularios_sugeridos = []
    arquivos = caminhos_resposta_gpt(caminho)
    for arquivo in arquivos:
        try:
            documento = ler_json(arquivo)
        except json.JSONDecodeError as erro:
            raise ValueError(f"JSON inválido em {arquivo}: {erro}") from erro
        if not isinstance(documento, dict) or documento.get("schema_version") != SCHEMA_RESPOSTA_GPT:
            raise ValueError(f"Schema GPT inválido em {arquivo}")
        trabalhos = documento.get("trabalhos")
        if not isinstance(trabalhos, list):
            raise ValueError(f"{arquivo}: trabalhos deve ser uma lista")
        resumo_documento = documento.get("resumo")
        if not isinstance(resumo_documento, dict) or resumo_documento.get("total_registros") != len(trabalhos):
            raise ValueError(f"{arquivo}: resumo.total_registros divergente")
        if not isinstance(documento.get("vocabularios_sugeridos"), dict):
            raise ValueError(f"{arquivo}: vocabularios_sugeridos deve ser objeto")
        if documento.get("vocabularios_sugeridos"):
            vocabularios_sugeridos.append(
                {"arquivo": str(arquivo), "sugestoes": documento["vocabularios_sugeridos"]}
            )
        for item in trabalhos:
            if not isinstance(item, dict):
                raise ValueError(f"{arquivo}: trabalho GPT não é objeto")
            identificador = limpar_texto(item.get("id"))
            if identificador not in esperados:
                raise ValueError(f"ID GPT inexistente: {identificador or '<vazio>'}")
            if identificador in sugestoes:
                raise ValueError(f"ID GPT duplicado: {identificador}")
            permitidas_item = {"id", "orientador_padronizado", "classificacoes"}
            extras = sorted(set(item) - permitidas_item)
            if extras:
                avisos.append(
                    f"{identificador}: campos proibidos/inesperados ignorados: {', '.join(extras)}"
                )

            linha_atual = linhas_por_id[identificador]
            orientador_informado = limpar_texto(item.get("orientador_padronizado"))
            if not orientador_informado:
                raise ValueError(f"{identificador}: orientador_padronizado vazio")
            orientador_canonico, erro_orientador, _ = normalizar_orientador(
                orientador_informado, vocabulario
            )
            if erro_orientador or orientador_canonico != linha_atual["Orientador"]:
                raise ValueError(
                    f"{identificador}: orientador_padronizado sugere pessoa diferente ou alias não confirmado"
                )

            classificacoes = item.get("classificacoes")
            if not isinstance(classificacoes, dict):
                raise ValueError(f"{identificador}: classificacoes ausente ou inválido")
            faltantes = sorted(set(CAMPOS_CLASSIFICACAO_GPT) - set(classificacoes))
            extras_classificacao = sorted(set(classificacoes) - set(CAMPOS_CLASSIFICACAO_GPT))
            if faltantes:
                raise ValueError(
                    f"{identificador}: resposta incompleta; faltam {', '.join(faltantes)}"
                )
            if extras_classificacao:
                raise ValueError(
                    f"{identificador}: classificações não permitidas: {', '.join(extras_classificacao)}"
                )

            classificacoes_validadas = {}
            for chave, coluna in CAMPOS_CLASSIFICACAO_GPT.items():
                detalhe = classificacoes[chave]
                if not isinstance(detalhe, dict):
                    raise ValueError(f"{identificador}/{chave}: detalhe inválido")
                confianca = normalizar_busca(detalhe.get("confianca"))
                if confianca not in CONFIANCAS_GPT:
                    raise ValueError(
                        f"{identificador}/{chave}: confiança inválida: {detalhe.get('confianca')!r}"
                    )
                if not isinstance(detalhe.get("revisao_manual"), bool):
                    raise ValueError(f"{identificador}/{chave}: revisao_manual deve ser booleano")
                evidencia = limpar_texto(detalhe.get("evidencia"))
                if not evidencia:
                    raise ValueError(f"{identificador}/{chave}: evidência vazia")
                valor_atual = limpar_texto(detalhe.get("valor_atual"))
                if not valor_atual:
                    raise ValueError(f"{identificador}/{chave}: valor_atual vazio")
                try:
                    valor_atual_canonico = canonizar_valor_gpt(
                        chave, valor_atual, vocabulario, permitidos
                    )
                except ValueError as erro:
                    raise ValueError(
                        f"{identificador}/{chave}: valor_atual diverge do projeto"
                    ) from erro
                if valor_atual_canonico != linha_atual[coluna]:
                    raise ValueError(
                        f"{identificador}/{chave}: valor_atual diverge do projeto"
                    )
                valor_informado = limpar_texto(detalhe.get("valor"))
                valor_canonico = canonizar_valor_gpt(
                    chave, valor_informado, vocabulario, permitidos
                )
                classificacoes_validadas[chave] = {
                    "valor_informado": valor_informado,
                    "valor_canonico": valor_canonico,
                    "confianca": confianca,
                    "evidencia": evidencia,
                    "motivo": limpar_texto(detalhe.get("motivo")),
                    "revisao_manual": detalhe["revisao_manual"],
                }

            linha_sugerida = classificacoes_validadas["linha_pesquisa"]["valor_canonico"]
            macro_sugerido = classificacoes_validadas["macroprojeto"]["valor_canonico"]
            if linha_sugerida != "A revisar" and macro_sugerido != "A revisar":
                codigo = macro_sugerido.split(" ", 1)[0]
                permitidos_linha = vocabulario.get("relacao_linha_macro", {}).get(
                    linha_sugerida, []
                )
                if codigo not in permitidos_linha:
                    raise ValueError(
                        f"{identificador}: Linha/Macro incompatíveis ({linha_sugerida} / {codigo})"
                    )

            sugestoes[identificador] = {
                "orientador_padronizado": orientador_canonico,
                "orientador_informado": orientador_informado,
                "classificacoes": classificacoes_validadas,
            }

    ausentes = sorted(esperados - set(sugestoes))
    if ausentes:
        amostra = ", ".join(ausentes[:12])
        sufixo = "..." if len(ausentes) > 12 else ""
        raise ValueError(f"Resposta GPT incompleta; IDs ausentes: {amostra}{sufixo}")
    return {
        "sugestoes": sugestoes,
        "avisos": avisos,
        "vocabularios_sugeridos": vocabularios_sugeridos,
        "arquivos": [str(item) for item in arquivos],
    }


def aplicar_sugestoes_gpt(linha: dict, sugestao: dict) -> tuple[list[dict], list[str]]:
    auditoria = []
    pendencias = []

    orientador_atual = linha["Orientador"]
    orientador_gpt = sugestao["orientador_padronizado"]
    acao_orientador = (
        "Aplicado automaticamente"
        if sugestao["orientador_informado"] != orientador_atual
        else "Mantido valor anterior"
    )
    auditoria.append(
        {
            "ID": linha["ID"],
            "Título": linha["Título da Dissertação"],
            "Campo": "Orientador",
            "Valor Atual": orientador_atual,
            "Valor GPT": sugestao["orientador_informado"],
            "Confiança": "alta",
            "Evidência": "Alias validado contra o mesmo orientador oficial.",
            "Ação": acao_orientador,
        }
    )
    linha["Orientador"] = orientador_gpt

    for chave, coluna in CAMPOS_CLASSIFICACAO_GPT.items():
        detalhe = sugestao["classificacoes"][chave]
        atual = linha[coluna]
        informado = detalhe["valor_informado"]
        canonico = detalhe["valor_canonico"]
        mesma_categoria = canonico == atual
        if detalhe["confianca"] == "baixa":
            acao = "Rejeitado"
        elif detalhe["confianca"] != "alta" or detalhe["revisao_manual"]:
            acao = "Revisão manual"
        elif mesma_categoria:
            acao = "Mantido valor anterior"
            linha[coluna] = canonico
        else:
            acao = "Aplicado automaticamente"
            linha[coluna] = canonico

        if chave == "macroprojeto" and acao in {"Aplicado automaticamente", "Mantido valor anterior"}:
            linha["Código do Macroprojeto"] = (
                canonico.split(" ", 1)[0] if canonico != "A revisar" else "A revisar"
            )

        if acao in {"Revisão manual", "Rejeitado"}:
            pendencias.append(f"sugestão GPT para {coluna}: {acao.lower()}")
        auditoria.append(
            {
                "ID": linha["ID"],
                "Título": linha["Título da Dissertação"],
                "Campo": coluna,
                "Valor Atual": atual,
                "Valor GPT": informado,
                "Confiança": detalhe["confianca"],
                "Evidência": detalhe["evidencia"],
                "Ação": acao,
            }
        )
    if pendencias:
        linha["Revisão Manual"] = "Sim"
        complemento = "Pendências GPT: " + "; ".join(pendencias)
        linha["Observações"] = (
            f"{linha['Observações']} | {complemento}" if linha["Observações"] else complemento
        )
    return auditoria, pendencias


def montar_relatorio(
    resumo: dict,
    linhas: list[dict],
    alteracoes: list[dict],
    revisoes: list[dict],
    antes: dict[str, Counter] | None = None,
    depois: dict[str, Counter] | None = None,
    revisao_gpt: list[dict] | None = None,
    avisos_gpt: list[str] | None = None,
) -> str:
    classificacoes_pendentes = {
        campo: sum(linha[campo] == "A revisar" for linha in linhas)
        for campo in ["Linha de Pesquisa", "Macroprojeto", "Base Epistemológica", "Nível de Aplicação", "Tipo de Produto Educacional", "Área Temática"]
    }
    links_produto: defaultdict[str, list[str]] = defaultdict(list)
    for linha in linhas:
        if linha["Link do Produto"]:
            links_produto[linha["Link do Produto"]].append(linha["ID"])
    produtos_repetidos = {url: ids for url, ids in links_produto.items() if len(ids) > 1}
    partes = [
        "RELATÓRIO DE VALIDAÇÃO — PAINEL EPT / LEME",
        f"Gerado em: {agora_utc()}",
        "",
        "RESUMO",
        f"Registros: {resumo['quantidade']}",
        f"IDs únicos: {resumo['ids_unicos']}",
        f"Vínculos confirmados: {resumo['vinculos_confirmados']}",
        f"Extrações de PDF com sucesso: {resumo['extracoes_sucesso']}",
        f"Erros estruturais: {resumo['erros_estruturais']}",
        f"Registros para revisão manual: {resumo['revisoes_manuais']}",
        f"Alterações determinísticas registradas: {len(alteracoes)}",
        "",
        "PADRONIZAÇÃO GLOBAL — ANTES E DEPOIS",
        "",
    ]
    antes = antes or {}
    depois = depois or {}
    for campo in ["Orientador", "Linha de Pesquisa", "Macroprojeto", *CAMPOS_VOCABULARIO_ABERTO, "Público-alvo"]:
        contador_antes = antes.get(campo, Counter())
        contador_depois = depois.get(campo, Counter())
        partes.append(f"{campo}: {len(contador_antes)} valores antes -> {len(contador_depois)} depois")
        principais = "; ".join(f"{valor} ({qtd})" for valor, qtd in contador_depois.most_common(12))
        partes.append(f"  Principais agrupamentos finais: {principais or 'nenhum'}")
    partes.extend(["", "CLASSIFICAÇÕES A REVISAR"])
    partes.extend(f"{campo}: {quantidade}" for campo, quantidade in classificacoes_pendentes.items())
    partes.extend(["", "DUPLICIDADES/CONFLITOS DE CONJUNTO"])
    partes.extend(resumo["duplicidades"] or ["Nenhuma duplicidade de PDF ou autor/título."])
    if produtos_repetidos:
        partes.append("Links de produto repetidos (não corrigidos automaticamente):")
        partes.extend(f"- {url}: {', '.join(ids)}" for url, ids in produtos_repetidos.items())
    if revisao_gpt is not None:
        contagem_acoes = Counter(item["Ação"] for item in revisao_gpt)
        partes.extend(["", "IMPORTAÇÃO GPT"])
        partes.append(f"Sugestões auditadas: {len(revisao_gpt)}")
        partes.extend(f"{acao}: {quantidade}" for acao, quantidade in contagem_acoes.items())
        partes.append(f"Alertas de campos ignorados: {len(avisos_gpt or [])}")
        partes.extend(f"- {aviso}" for aviso in (avisos_gpt or []))
    partes.extend(["", "REGISTROS PARA REVISÃO"])
    partes.extend(f"- {item['ID']}: {item['Motivos']}" for item in revisoes)
    if not revisoes:
        partes.append("Nenhum.")
    return "\n".join(partes) + "\n"


def bool_planilha(valor) -> bool:
    return normalizar_busca(valor) in {"true", "sim", "1", "yes"}


def carregar_registros_entrada() -> tuple[dict, list[dict]]:
    """Recompõe a visão auditável usando o JSON para GPT e a planilha oficial."""
    entrada = ler_json(ARQUIVO_ENTRADA)
    if not entrada or entrada.get("schema_version") != SCHEMA_ENTRADA_GPT:
        raise ValueError(
            f"Entrada GPT inválida: {ARQUIVO_ENTRADA}. Execute primeiro 2_PdfParaJson.py"
        )
    if not ARQUIVO_DADOS_OFICIAIS.is_file():
        raise FileNotFoundError(f"Planilha oficial ausente: {ARQUIVO_DADOS_OFICIAIS}")
    quadro = pd.read_excel(ARQUIVO_DADOS_OFICIAIS, dtype=str).fillna("")
    oficiais = {limpar_texto(linha["ID"]): linha for _, linha in quadro.iterrows()}
    trabalhos = entrada.get("trabalhos")
    if not isinstance(trabalhos, list):
        raise ValueError("para_gpt.json não contém uma lista de trabalhos")
    ids_entrada = [limpar_texto(item.get("id")) for item in trabalhos]
    if len(ids_entrada) != len(set(ids_entrada)):
        raise ValueError("para_gpt.json contém IDs duplicados")
    if set(ids_entrada) != set(oficiais):
        faltantes = sorted(set(oficiais) - set(ids_entrada))
        extras = sorted(set(ids_entrada) - set(oficiais))
        raise ValueError(
            f"IDs divergentes entre para_gpt.json e planilha oficial; faltantes={faltantes}, extras={extras}"
        )

    mapa_oficial = {
        "ano": "Ano",
        "titulo": "Título da Dissertação",
        "autor": "Autor",
        "instituicao": "Instituição",
        "campus": "Campus",
        "programa": "Programa",
    }
    mapa_classificacao = {
        "linha_pesquisa": "Linha de Pesquisa",
        "macroprojeto": "Macroprojeto",
        "base_epistemologica": "Base Epistemológica",
        "nivel_aplicacao": "Nível de Aplicação",
        "tipo_produto": "Tipo de Produto Educacional",
        "area_tematica": "Área Temática",
        "publico_alvo": "Público-alvo",
    }
    registros = []
    for trabalho in trabalhos:
        identificador = limpar_texto(trabalho.get("id"))
        oficial = oficiais[identificador]
        for campo_json, coluna in mapa_oficial.items():
            if normalizar_busca(trabalho.get(campo_json)) != normalizar_busca(oficial.get(coluna)):
                raise ValueError(
                    f"{identificador}: dado oficial divergente em {coluna} entre as etapas 1 e 2"
                )
        classificacao = trabalho.get("classificacao_atual")
        if not isinstance(classificacao, dict):
            raise ValueError(f"{identificador}: classificacao_atual ausente")
        valores = {}
        valores_originais = {}
        for chave, coluna in mapa_classificacao.items():
            detalhe = classificacao.get(chave)
            if not isinstance(detalhe, dict):
                raise ValueError(f"{identificador}: classificação ausente: {chave}")
            valores_originais[coluna] = limpar_texto(detalhe.get("valor_original")) or "A revisar"
            valores[coluna] = limpar_texto(detalhe.get("valor_padronizado")) or "A revisar"
        codigo = re.search(r"\bMP\s*([1-6])\b", valores["Macroprojeto"], flags=re.I)
        valores["Código do Macroprojeto"] = f"MP{codigo.group(1)}" if codigo else "A revisar"
        valores["Revisão Manual"] = "Sim" if "A revisar" in valores.values() else "Não"
        valores["Observações"] = (
            "Classificação sem evidência suficiente; revisar manualmente."
            if "A revisar" in valores.values()
            else ""
        )
        conteudo = {
            "resumo": trabalho.get("resumo"),
            "palavras_chave": trabalho.get("palavras_chave") or [],
            "objetivos": trabalho.get("objetivo_geral"),
            "metodologia": trabalho.get("metodologia"),
            "produto_educacional_extraido": trabalho.get("descricao_produto"),
            "publico_alvo_extraido": trabalho.get("publico_alvo_extraido"),
            "base_teorica_extraida": trabalho.get("base_teorica_extraida"),
        }
        evidencia_ocr = any(
            limpar_texto(conteudo.get(campo))
            for campo in ["resumo", "objetivos", "metodologia", "produto_educacional_extraido"]
        ) or bool(conteudo["palavras_chave"])
        conflitos = []
        texto_conflitos = limpar_texto(oficial.get("Conflitos de Fonte"))
        if texto_conflitos:
            try:
                conflitos = json.loads(texto_conflitos)
            except json.JSONDecodeError:
                conflitos = [{"campo": "fontes", "detalhe": texto_conflitos}]
        registros.append(
            {
                "id": identificador,
                "ano": limpar_texto(oficial.get("Ano")),
                "titulo": limpar_texto(oficial.get("Título da Dissertação")),
                "autor": limpar_texto(oficial.get("Autor")),
                "orientador": limpar_texto(oficial.get("Orientador")),
                "coorientador": limpar_texto(oficial.get("Coorientador")),
                "produto_educacional": limpar_texto(oficial.get("Produto Educacional")),
                "instituicao": limpar_texto(oficial.get("Instituição")),
                "campus": limpar_texto(oficial.get("Campus")),
                "programa": limpar_texto(oficial.get("Programa")),
                "instituicao_campus": f"IFC {limpar_texto(oficial.get('Campus'))}",
                "url_pdf": limpar_texto(oficial.get("Link do PDF")),
                "url_produto": limpar_texto(oficial.get("Link do Produto")),
                "arquivo_pdf": {
                    "caminho": str(Path("pdf") / limpar_texto(oficial.get("Nome do PDF"))),
                    "sha256": limpar_texto(oficial.get("SHA-256")),
                },
                "conteudo_pdf": conteudo,
                "classificacoes": {
                    "valores": valores,
                    "valores_originais": valores_originais,
                    "origem": "para_gpt.json",
                },
                "extracao_pdf": {"status": "sucesso" if evidencia_ocr else "sem_evidencia_extraida"},
                "validacao": {
                    "vinculo_confirmado": bool_planilha(oficial.get("Vínculo Confirmado")),
                    "conflitos": conflitos,
                },
            }
        )
    return entrada, registros


def executar(classificacoes_gpt: Path = ARQUIVO_RESPOSTA_GPT) -> dict:
    consolidado, registros = carregar_registros_entrada()
    vocabulario = ler_json(ARQUIVO_VOCABULARIOS)
    if not vocabulario or vocabulario.get("schema_version") != SCHEMA_VOCABULARIOS:
        raise ValueError(f"Vocabulário inválido: {ARQUIVO_VOCABULARIOS}")
    if not classificacoes_gpt.is_file():
        raise FileNotFoundError(
            f"Resposta GPT ausente: {classificacoes_gpt}. Salve o JSON único nesse caminho."
        )
    registros = sorted(registros, key=lambda item: item.get("id") or "ZZZ")
    lexico = construir_lexico_textual(registros)
    antes = inventario_antes(registros)
    linhas = []
    alteracoes = []
    for registro in registros:
        linha, mudancas, _ = construir_linha(registro, vocabulario, lexico)
        linhas.append(linha)
        alteracoes.extend(mudancas)
    revisao_gpt: list[dict] = []
    avisos_gpt: list[str] = []
    if classificacoes_gpt:
        importacao = validar_respostas_gpt(classificacoes_gpt, registros, linhas, vocabulario)
        avisos_gpt = importacao["avisos"]
        sugestoes = importacao["sugestoes"]
        for linha in linhas:
            auditoria, _ = aplicar_sugestoes_gpt(linha, sugestoes[linha["ID"]])
            revisao_gpt.extend(auditoria)
        for grupo in importacao["vocabularios_sugeridos"]:
            revisao_gpt.append(
                {
                    "ID": "GLOBAL",
                    "Título": "Sugestão de vocabulário",
                    "Campo": "Vocabulários controlados",
                    "Valor Atual": "Vocabulário versionado preservado",
                    "Valor GPT": limitar_contexto_gpt(
                        json.dumps(grupo["sugestoes"], ensure_ascii=False), 3000
                    ),
                    "Confiança": "A revisar",
                    "Evidência": f"Arquivo: {grupo['arquivo']}",
                    "Ação": "Revisão manual",
                }
            )
        for aviso in avisos_gpt:
            revisao_gpt.append(
                {
                    "ID": "GLOBAL",
                    "Título": "Campo proibido ou inesperado",
                    "Campo": "Importação GPT",
                    "Valor Atual": "Dados oficiais preservados",
                    "Valor GPT": aviso,
                    "Confiança": "Não aplicável",
                    "Evidência": "Validador de campos permitidos",
                    "Ação": "Rejeitado",
                }
            )
    depois = inventario_depois(linhas)
    aliases_orientador = tabela_aliases_orientador(registros, vocabulario)
    validacoes, revisoes, resumo = validar_conjunto(linhas)
    resumo["categorias_antes_depois"] = {
        campo: {"antes": len(antes[campo]), "depois": len(depois[campo])}
        for campo in depois
    }
    resumo["importacao_gpt"] = {
        "utilizada": bool(classificacoes_gpt),
        "sugestoes_auditadas": len(revisao_gpt),
        "avisos": len(avisos_gpt),
        "acoes": dict(Counter(item["Ação"] for item in revisao_gpt)),
    }
    salvar_planilha(
        linhas, alteracoes, validacoes, revisoes, aliases_orientador, revisao_gpt
    )
    salvar_texto_atomico(
        montar_relatorio(
            resumo,
            linhas,
            alteracoes,
            revisoes,
            antes,
            depois,
            revisao_gpt if classificacoes_gpt else None,
            avisos_gpt,
        ),
        ARQUIVO_RELATORIO,
    )
    return resumo


def main() -> int:
    args = argumentos()
    try:
        resumo = executar(args.classificacoes_gpt)
    except (FileNotFoundError, ValueError, json.JSONDecodeError) as erro:
        print(f"ERRO: {erro}")
        return 2
    print(json.dumps(resumo, ensure_ascii=False, indent=2))
    rotulos = {
        "Orientador": "Orientadores distintos",
        "Nível de Aplicação": "Níveis distintos",
        "Base Epistemológica": "Bases Epistemológicas distintas",
        "Tipo de Produto Educacional": "Tipos de Produto distintos",
        "Área Temática": "Áreas Temáticas distintas",
        "Público-alvo": "Públicos-alvo distintos",
    }
    print("\nPADRONIZAÇÃO — ANTES E DEPOIS")
    for campo, rotulo in rotulos.items():
        contagens = resumo["categorias_antes_depois"].get(campo, {})
        print(f"{rotulo} antes: {contagens.get('antes', 0)}")
        print(f"{rotulo} depois: {contagens.get('depois', 0)}")
    return 0 if resumo["erros_estruturais"] == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
